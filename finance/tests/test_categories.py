"""Categories: unique names and the one-level subcategory rule, types, display names and
grouping, the category list (search, usage, mass re-categorize, delete), the category
dropdowns in forms, and the starter categories."""

import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest import mock

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.urls import reverse
from openpyxl import Workbook

from .. import starter_categories
from ..forms import CategoryForm, TransactionForm
from ..models import Account, Category, Payee, Transaction, grouped_categories
from ..services.importing import _resolve_category
from .base import AlexsStashTestCase, LoggedInTestCase, LoggedInWithAccountTestCase, StashTestCase

Type = Category.CategoryType

EXPENSE = Category.CategoryType.EXPENSE


class IsActiveRemovedTests(StashTestCase):
    def test_category_has_no_is_active_field(self):
        self.assertFalse(hasattr(Category(name="x"), "is_active"))

    def test_category_form_has_no_is_active_field(self):
        self.assertNotIn("is_active", CategoryForm().fields)


class DisplayNameTests(StashTestCase):
    def test_full_name_shows_parent_then_child_with_an_arrow(self):
        food = Category.objects.create(name="Food", category_type=Type.EXPENSE)
        groceries = Category.objects.create(name="Groceries", parent=food, category_type=Type.EXPENSE)
        self.assertEqual(groceries.full_name, "Food > Groceries")
        self.assertEqual(food.full_name, "Food")

    def test_transfer_category_is_named_in_brackets(self):
        account = Account.objects.create(name="Cash")
        self.assertEqual(account.transfer_category.name, "[Cash]")
        account.name = "Wallet"
        account.save()
        account.transfer_category.refresh_from_db()
        self.assertEqual(account.transfer_category.name, "[Wallet]")


class OneLevelOfSubcategoriesTests(StashTestCase):
    def test_a_subcategory_is_not_offered_as_a_parent(self):
        food = Category.objects.create(name="Food", category_type=Type.EXPENSE)
        groceries = Category.objects.create(name="Groceries", parent=food, category_type=Type.EXPENSE)
        self.assertNotIn(groceries, CategoryForm().fields["parent"].queryset)
        self.assertIn(food, CategoryForm().fields["parent"].queryset)

    def test_a_category_cannot_become_a_subcategory_of_itself(self):
        food = Category.objects.create(name="Food", category_type=Type.EXPENSE)
        form = CategoryForm({"name": "Food", "parent": food.pk, "category_type": "expense"}, instance=food)
        self.assertFalse(form.is_valid())
        self.assertIn("parent", form.errors)

    def test_a_category_with_children_cannot_be_demoted_into_a_subcategory(self):
        food = Category.objects.create(name="Food", category_type=Type.EXPENSE)
        Category.objects.create(name="Groceries", parent=food, category_type=Type.EXPENSE)
        rent = Category.objects.create(name="Rent", category_type=Type.EXPENSE)
        form = CategoryForm({"name": "Food", "parent": rent.pk, "category_type": "expense"}, instance=food)
        self.assertFalse(form.is_valid())
        self.assertIn("parent", form.errors)

    def test_a_childless_category_can_become_a_subcategory(self):
        rent = Category.objects.create(name="Rent", category_type=Type.EXPENSE)
        bills = Category.objects.create(name="Bills", category_type=Type.EXPENSE)
        form = CategoryForm({"name": "Rent", "parent": bills.pk, "category_type": "expense"}, instance=rent)
        self.assertTrue(form.is_valid(), form.errors)


class GroupedCategoriesTests(StashTestCase):
    def test_income_then_expense_then_transfer_each_sorted_by_parent_then_name(self):
        Category.objects.create(name="Salary", category_type=Type.INCOME)
        Category.objects.create(name="Bonus", category_type=Type.INCOME)
        food = Category.objects.create(name="Food", category_type=Type.EXPENSE)
        Category.objects.create(name="Groceries", parent=food, category_type=Type.EXPENSE)
        Category.objects.create(name="Rent", category_type=Type.EXPENSE)
        account = Account.objects.create(name="Checking")  # brings its own Transfer category

        groups = grouped_categories(Category.objects.select_related("parent").all())

        self.assertEqual([g["key"] for g in groups], ["income", "expense", "transfer"])
        self.assertEqual([c.name for c in groups[0]["categories"]], ["Bonus", "Salary"])
        self.assertEqual([c.name for c in groups[1]["categories"]], ["Food", "Groceries", "Rent"])
        self.assertEqual([c.name for c in groups[2]["categories"]], [account.transfer_category.name])

    def test_a_category_type_that_differs_from_data_still_groups_by_its_own_type(self):
        # Grouping reads each row's own category_type - it doesn't try to inherit a parent's.
        parent = Category.objects.create(name="Parent", category_type=Type.INCOME)
        child = Category.objects.create(name="Child", parent=parent, category_type=Type.EXPENSE)
        groups = {g["key"]: [c.name for c in g["categories"]] for g in grouped_categories([parent, child])}
        self.assertEqual(groups, {"income": ["Parent"], "expense": ["Child"]})


class CategoryParentTypeFilterTests(StashTestCase):
    def test_parent_options_are_tagged_with_their_own_category_type(self):
        food = Category.objects.create(name="Food", category_type=Type.EXPENSE)
        salary = Category.objects.create(name="Salary", category_type=Type.INCOME)
        form = CategoryForm()
        widget = form.fields["parent"].widget
        self.assertEqual(widget.categories_by_id[food.pk].category_type, Type.EXPENSE)
        self.assertEqual(widget.categories_by_id[salary.pk].category_type, Type.INCOME)

    def test_a_transfer_category_is_not_offered_as_a_parent(self):
        account = Account.objects.create(name="Checking")
        form = CategoryForm()
        self.assertNotIn(account.transfer_category.pk, form.fields["parent"].widget.categories_by_id)
        self.assertNotIn(account.transfer_category, form.fields["parent"].queryset)


class CategorySearchTests(LoggedInTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.food = Category.objects.create(name="Food", category_type=Type.EXPENSE)
        Category.objects.create(name="Groceries", parent=cls.food, category_type=Type.EXPENSE)
        Category.objects.create(name="Rent", category_type=Type.EXPENSE)

    def test_search_matches_by_own_name_or_parent_name(self):
        response = self.client.get(reverse("finance:category_list"), {"q": "food"})
        names = {c.name for c in response.context["categories"]}
        self.assertEqual(names, {"Food", "Groceries"})

    def test_search_with_no_matches_shows_a_message(self):
        response = self.client.get(reverse("finance:category_list"), {"q": "nonexistent"})
        self.assertContains(response, "No categories match")


class CategoryUsageAndRecategorizeTests(LoggedInTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.account = Account.objects.create(name="Checking")
        cls.groceries = Category.objects.create(name="Groceries", category_type=Type.EXPENSE)
        cls.household = Category.objects.create(name="Household", category_type=Type.EXPENSE)
        cls.txns = [
            Transaction.objects.create(
                account=cls.account, date=date(2026, 1, d), category=cls.groceries, amount=Decimal("-10.00")
            )
            for d in (1, 2, 3)
        ]

    def test_usage_count_shown_on_the_category_list(self):
        response = self.client.get(reverse("finance:category_list"))
        categories = {c.pk: c.usage_count for group in response.context["groups"] for c in group["categories"]}
        self.assertEqual(categories[self.groceries.pk], 3)
        self.assertEqual(categories[self.household.pk], 0)

    def test_usage_view_lists_only_that_categorys_own_transactions(self):
        sub = Category.objects.create(name="Organic", parent=self.groceries, category_type=Type.EXPENSE)
        Transaction.objects.create(account=self.account, date=date(2026, 1, 4), category=sub, amount=Decimal("-5.00"))
        response = self.client.get(reverse("finance:category_usage", args=[self.groceries.pk]))
        self.assertEqual(len(response.context["transactions"]), 3)

    def test_moving_selected_transactions_updates_only_those_selected(self):
        response = self.client.post(
            reverse("finance:category_usage", args=[self.groceries.pk]),
            {
                "transaction_id": [self.txns[0].pk, self.txns[1].pk],
                "move_to": self.household.pk,
            },
        )
        self.assertEqual(response.status_code, 302)
        self.txns[0].refresh_from_db()
        self.txns[1].refresh_from_db()
        self.txns[2].refresh_from_db()
        self.assertEqual(self.txns[0].category, self.household)
        self.assertEqual(self.txns[1].category, self.household)
        self.assertEqual(self.txns[2].category, self.groceries)

    def test_cannot_move_transactions_into_a_transfer_category(self):
        self.client.post(
            reverse("finance:category_usage", args=[self.groceries.pk]),
            {
                "transaction_id": [self.txns[0].pk],
                "move_to": self.account.transfer_category.pk,
            },
        )
        self.txns[0].refresh_from_db()
        self.assertEqual(self.txns[0].category, self.groceries)

    def test_selecting_nothing_shows_an_error_and_moves_nothing(self):
        self.client.post(reverse("finance:category_usage", args=[self.groceries.pk]), {"move_to": self.household.pk})
        self.txns[0].refresh_from_db()
        self.assertEqual(self.txns[0].category, self.groceries)


class SystemManagedCategoryStillGuardedTests(LoggedInTestCase):
    def test_transfer_category_not_offered_in_its_own_usage_views_move_to_list(self):
        account = Account.objects.create(name="Checking")
        other = Account.objects.create(name="Savings")
        response = self.client.get(reverse("finance:category_usage", args=[account.transfer_category.pk]))
        self.assertNotIn(other.transfer_category, response.context["other_categories"])


class CategoryUniquenessTests(StashTestCase):
    """2.3 (F3): top-level names are unique too, and both rules ignore case."""

    def test_the_database_rejects_a_second_top_level_category_in_another_case(self):
        Category.objects.create(name="Food", category_type=EXPENSE)
        with self.assertRaises(IntegrityError), transaction.atomic(using=settings.TEST_STASH):
            Category.objects.create(name="FOOD", category_type=EXPENSE)

    def test_the_database_rejects_a_case_duplicate_subcategory(self):
        food = Category.objects.create(name="Food", category_type=EXPENSE)
        Category.objects.create(name="Snacks", parent=food, category_type=EXPENSE)
        with self.assertRaises(IntegrityError), transaction.atomic(using=settings.TEST_STASH):
            Category.objects.create(name="snacks", parent=food, category_type=EXPENSE)

    def test_the_same_subcategory_name_under_two_parents_is_fine(self):
        food = Category.objects.create(name="Food", category_type=EXPENSE)
        home = Category.objects.create(name="Home", category_type=EXPENSE)
        Category.objects.create(name="Other", parent=food, category_type=EXPENSE)
        Category.objects.create(name="Other", parent=home, category_type=EXPENSE)

    def test_two_accounts_with_one_name_still_get_their_transfer_categories(self):
        Account.objects.create(name="Cash")
        Account.objects.create(name="Cash")
        self.assertEqual(Category.objects.filter(name="[Cash]").count(), 2)

    def test_the_form_rejects_a_duplicate_with_a_clear_message(self):
        Category.objects.create(name="Food", category_type=EXPENSE)
        form = CategoryForm(data={"name": "food", "category_type": EXPENSE})
        self.assertFalse(form.is_valid())
        self.assertEqual(form.errors["name"], ["A category with this name already exists."])

    def test_the_form_rejects_a_duplicate_subcategory(self):
        food = Category.objects.create(name="Food", category_type=EXPENSE)
        Category.objects.create(name="Snacks", parent=food, category_type=EXPENSE)
        form = CategoryForm(data={"name": "SNACKS", "parent": food.pk, "category_type": EXPENSE})
        self.assertFalse(form.is_valid())
        self.assertEqual(form.errors["name"], ["'Food' already has a subcategory with this name."])


class OneLevelRuleTests(StashTestCase):
    """2.6 (F13): the one-level subcategory rule lives on the model, so the admin keeps it too."""

    @classmethod
    def setUpTestData(cls):
        cls.food = Category.objects.create(name="Food", category_type=EXPENSE)
        cls.snacks = Category.objects.create(name="Snacks", parent=cls.food, category_type=EXPENSE)

    def assert_parent_error(self, category, message):
        with self.assertRaises(ValidationError) as caught:
            category.full_clean()
        self.assertEqual(caught.exception.message_dict["parent"], [message])

    def test_a_subcategory_cant_have_subcategories(self):
        self.assert_parent_error(
            Category(name="Crisps", parent=self.snacks, category_type=EXPENSE),
            "A subcategory can't have its own subcategories - pick a main category.",
        )

    def test_a_category_with_subcategories_cant_become_one(self):
        home = Category.objects.create(name="Home", category_type=EXPENSE)
        self.food.parent = home
        self.assert_parent_error(
            self.food,
            "This category already has its own subcategories, so it can't become one itself "
            "— move or remove them first.",
        )

    def test_a_category_cant_sit_under_itself(self):
        self.food.subcategories.all().delete()
        self.food.parent = self.food
        self.assert_parent_error(self.food, "A category can't be placed under itself.")

    def test_a_transfer_category_cant_be_a_parent(self):
        transfer = Account.objects.create(name="Cash").transfer_category
        self.assert_parent_error(
            Category(name="Coins", parent=transfer, category_type=EXPENSE),
            "An account's transfer category can't have subcategories.",
        )

    def test_the_category_form_shows_the_models_message(self):
        home = Category.objects.create(name="Home", category_type=EXPENSE)
        form = CategoryForm(data={"name": "Food", "parent": home.pk, "category_type": EXPENSE}, instance=self.food)
        self.assertFalse(form.is_valid())
        self.assertIn("already has its own subcategories", form.errors["parent"][0])


class CategoryLoopTests(StashTestCase):
    @classmethod
    def setUpTestData(cls):
        cls.parent = Category.objects.create(name="Parent")
        cls.child = Category.objects.create(name="Child", parent=cls.parent)

    def test_descendants_stops_at_a_loop(self):
        Category.objects.filter(pk=self.parent.pk).update(parent=self.child)
        self.parent.refresh_from_db()
        self.assertEqual([c.pk for c in self.parent.descendants()], [self.child.pk])

    def test_form_rejects_itself_or_its_own_subcategory_as_parent(self):
        for invalid_parent in (self.parent, self.child):
            form = CategoryForm(
                {"name": "Parent", "parent": invalid_parent.pk, "category_type": "expense"},
                instance=self.parent,
            )
            self.assertFalse(form.is_valid())
            self.assertIn("parent", form.errors)


class CategoryDeleteTests(LoggedInWithAccountTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.food = Category.objects.create(name="Food")
        cls.groceries = Category.objects.create(name="Groceries", parent=cls.food)
        cls.organic = Category.objects.create(name="Organic", parent=cls.groceries)
        cls.txn = Transaction.objects.create(
            account=cls.account,
            date=date(2026, 1, 5),
            category=cls.organic,
            amount=Decimal("-5.00"),
        )

    def test_confirmation_lists_every_subcategory_that_would_be_deleted(self):
        response = self.client.get(reverse("finance:category_delete", args=[self.food.pk]))
        self.assertContains(response, "also delete 2 subcategories")
        self.assertContains(response, "Food &gt; Groceries")
        self.assertContains(response, "Groceries &gt; Organic")
        self.assertContains(response, "1 transaction in")

    def test_leaf_category_shows_no_subcategory_warning(self):
        response = self.client.get(reverse("finance:category_delete", args=[self.organic.pk]))
        self.assertNotContains(response, "subcategor")

    def test_deleting_keeps_transactions_uncategorized(self):
        self.client.post(reverse("finance:category_delete", args=[self.food.pk]))
        # The account's own transfer category (Phase 3) is unrelated and stays.
        self.assertFalse(Category.objects.filter(account__isnull=True).exists())
        self.txn.refresh_from_db()
        self.assertIsNone(self.txn.category)


class DeleteConfirmationRedesignTests(LoggedInWithAccountTestCase):
    def test_shows_an_icon_and_a_clear_warning(self):
        Category.objects.create(name="Groceries", category_type=Category.CategoryType.EXPENSE)
        cat = Category.objects.get(name="Groceries")
        response = self.client.get(reverse("finance:category_delete", args=[cat.pk]))
        self.assertContains(response, "stat-icon-danger")
        self.assertContains(response, "This can't be undone.")
        self.assertContains(response, "Yes, delete it")


class UsageEmptyStylingTests(LoggedInWithAccountTestCase):
    def test_a_category_with_no_transactions_is_styled_pale(self):
        Category.objects.create(name="Unused", category_type=Category.CategoryType.EXPENSE)
        response = self.client.get(reverse("finance:category_list"))
        self.assertContains(response, "usage-empty")

    def test_a_category_with_transactions_is_not_styled_pale(self):
        used = Category.objects.create(name="Used", category_type=Category.CategoryType.EXPENSE)
        Transaction.objects.create(account=self.account, date=date(2026, 1, 1), category=used, amount=Decimal("-1.00"))
        response = self.client.get(reverse("finance:category_list"))
        # The account's own transfer category always has 0 usage, so some "usage-empty" will
        # exist regardless - check the specific "Used" row's link isn't the pale class.
        content = response.content.decode()
        used_index = content.index("Used<")
        nearby = content[used_index : used_index + 400]
        self.assertNotIn("usage-empty", nearby)

    def test_a_payee_with_no_transactions_is_styled_pale(self):
        Payee.objects.create(name="Ghost Payee")
        response = self.client.get(reverse("finance:payee_list"))
        self.assertContains(response, "usage-empty")


class CategoryDropdownOrderingTests(LoggedInWithAccountTestCase):
    def test_subcategories_sort_next_to_their_own_parent(self):
        food = Category.objects.create(name="Food", category_type=Category.CategoryType.EXPENSE)
        Category.objects.create(name="Groceries", parent=food, category_type=Category.CategoryType.EXPENSE)
        personal = Category.objects.create(name="Personal Care", category_type=Category.CategoryType.EXPENSE)
        Category.objects.create(name="Haircut", parent=personal, category_type=Category.CategoryType.EXPENSE)
        form = TransactionForm()
        names = [str(c) for c in form.fields["category"].queryset]
        # "Personal Care" and its own "Haircut" subcategory land next to each other, not split
        # into a top-level block and a separate subcategories block.
        personal_index = names.index("Personal Care")
        self.assertEqual(names[personal_index + 1], "Personal Care > Haircut")


class IndentedCategoryDropdownTests(LoggedInWithAccountTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        food = Category.objects.create(name="Food", category_type=Category.CategoryType.EXPENSE)
        Category.objects.create(name="Groceries", parent=food, category_type=Category.CategoryType.EXPENSE)

    def _assert_indented(self, html):
        self.assertIn(">Food</option>", html)
        self.assertIn(">    Groceries</option>", html.replace("&nbsp;", " "))
        self.assertNotIn("Food &gt; Groceries", html)

    def test_transaction_form(self):
        html = self.client.get(reverse("finance:transaction_add")).content.decode()
        self._assert_indented(html)
        self.assertLess(html.index(">Food</option>"), html.index("Groceries</option>"))

    def test_reminder_form(self):
        self._assert_indented(self.client.get(reverse("finance:reminder_add")).content.decode())


class NewSubcategoryTypeTests(StashTestCase):
    """6.9 (F38): a new subcategory takes its main category's type, wherever it's created."""

    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.path = Path(tmp.name) / "starter_categories.xlsx"
        patcher = mock.patch.object(starter_categories, "starter_categories_path", return_value=self.path)
        patcher.start()
        self.addCleanup(patcher.stop)

    def _write(self, rows):

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Extended"
        sheet.append(["Type", "Category", "Subcategory"])
        for row in rows:
            sheet.append(row)
        workbook.save(self.path)

    def test_starter_subcategory_follows_its_main_category_not_its_row(self):
        self._write([("Expense", "Gifts", None), ("Income", "Gifts", "Birthday money")])
        starter_categories.add_starter_categories("Extended")
        child = Category.objects.get(name="Birthday money")
        self.assertEqual(child.category_type, Category.CategoryType.EXPENSE)

    def test_starter_never_puts_a_subcategory_under_a_transfer_category(self):
        account = Account.objects.create(name="Wallet")
        self._write([("Expense", "[Wallet]", "Snacks")])
        starter_categories.add_starter_categories("Extended")
        self.assertFalse(account.transfer_category.subcategories.exists())
        main = Category.objects.main_categories().get(name="[Wallet]")
        self.assertEqual(main.subcategories.get().name, "Snacks")

    def test_importer_uses_the_same_rule(self):
        Category.objects.create(name="Gifts", category_type=Category.CategoryType.EXPENSE)
        child = _resolve_category("Gifts:Birthday money", amount=50)
        self.assertEqual(child.category_type, Category.CategoryType.EXPENSE)
        self.assertTrue(child.is_auto_generated)


class StarterCategoriesXlsxTests(StashTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "starter_categories.xlsx"
        self.patcher = mock.patch("finance.starter_categories.starter_categories_path", return_value=self.path)
        self.patcher.start()
        self.addCleanup(self.patcher.stop)

    def _write(self, sheets):

        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, rows in sheets.items():
            sheet = workbook.create_sheet(name)
            sheet.append(["Type", "Category", "Subcategory"])
            for row in rows:
                sheet.append(row)
        workbook.save(self.path)

    def test_missing_file_has_no_available_sheets(self):
        self.assertEqual(starter_categories.available_sheets(), [])

    def test_available_sheets_matches_whats_in_the_file(self):
        self._write({"Basic": [("Expense", "Groceries", None)]})
        self.assertEqual(starter_categories.available_sheets(), ["Basic"])

    def test_basic_creates_top_level_categories_only(self):
        self._write({"Basic": [("Expense", "Groceries", None), ("Income", "Work Income", None)]})
        created = starter_categories.add_starter_categories("Basic")
        self.assertEqual(created, ["Groceries", "Work Income"])
        self.assertEqual(Category.objects.get(name="Groceries").category_type, Category.CategoryType.EXPENSE)
        self.assertEqual(Category.objects.get(name="Work Income").category_type, Category.CategoryType.INCOME)
        self.assertFalse(Category.objects.filter(parent__isnull=False).exists())

    def test_extended_creates_subcategories_under_a_shared_parent(self):
        self._write(
            {
                "Extended": [
                    ("Expense", "Auto", "Fuel"),
                    ("Expense", "Auto", "Repairs"),
                    ("Income", "Other Income", "Gifts"),
                ]
            }
        )
        created = starter_categories.add_starter_categories("Extended")
        auto = Category.objects.get(name="Auto", parent__isnull=True)
        self.assertEqual(set(auto.subcategories.values_list("name", flat=True)), {"Fuel", "Repairs"})
        self.assertEqual(created, ["Auto", "Auto: Fuel", "Auto: Repairs", "Other Income", "Other Income: Gifts"])

    def test_running_twice_creates_nothing_the_second_time(self):
        self._write({"Basic": [("Expense", "Groceries", None)]})
        starter_categories.add_starter_categories("Basic")
        second_run = starter_categories.add_starter_categories("Basic")
        self.assertEqual(second_run, [])

    def test_existing_category_is_matched_case_insensitively(self):
        Category.objects.create(name="groceries", category_type=Category.CategoryType.EXPENSE)
        self._write({"Basic": [("Expense", "Groceries", None)]})
        created = starter_categories.add_starter_categories("Basic")
        self.assertEqual(created, [])
        self.assertEqual(Category.objects.filter(name__iexact="groceries").count(), 1)

    def test_unavailable_sheet_returns_none(self):
        self._write({"Basic": [("Expense", "Groceries", None)]})
        self.assertIsNone(starter_categories.add_starter_categories("Extended"))


class OldStarterCategoriesButtonIsGoneTests(AlexsStashTestCase):
    def test_category_list_has_no_add_starter_categories_button(self):
        response = self.client.get(reverse("finance:category_list"))
        self.assertNotContains(response, "Add starter categories")
