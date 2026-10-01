"""Starter categories (Phase 15 / L5): `starter_categories.xlsx` at the project root, two
sheets - Basic and Extended - each with Type | Category | Subcategory columns. Read only by
the welcome wizard's Categories step; there's no button to run this again afterwards - delete
what you don't want instead.
"""
from pathlib import Path

from django.conf import settings

from .models import Category

STARTER_CATEGORIES_FILENAME = "starter_categories.xlsx"
SHEET_CHOICES = ["Basic", "Extended"]


def starter_categories_path():
    return Path(settings.BASE_DIR) / STARTER_CATEGORIES_FILENAME


def available_sheets():
    """Which of Basic/Extended can actually be offered right now - [] if the file itself is
    missing, so the welcome step can disable choices cleanly instead of erroring."""
    path = starter_categories_path()
    if not path.exists():
        return []
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    try:
        names = set(workbook.sheetnames)
    finally:
        workbook.close()
    return [sheet for sheet in SHEET_CHOICES if sheet in names]


def _read_rows(sheet_name):
    """(type, category, subcategory_or_None) for every non-blank row, in file order."""
    from openpyxl import load_workbook

    workbook = load_workbook(starter_categories_path(), read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows = []
        for category_type, category, subcategory in sheet.iter_rows(min_row=2, values_only=True):
            if not category_type or not category:
                continue
            rows.append((
                str(category_type).strip(),
                str(category).strip(),
                str(subcategory).strip() if subcategory not in (None, "") else None,
            ))
        return rows
    finally:
        workbook.close()


def add_starter_categories(sheet_name):
    """Creates every category (and subcategory) from the given sheet that doesn't already
    exist (matched by name, case-insensitive, among top-level categories for a parent, or
    among that parent's own children for a subcategory) - safe to run more than once. Returns
    the list of names actually created, or None if the sheet isn't available."""
    if sheet_name not in available_sheets():
        return None

    top_level = {
        category.name.lower(): category
        for category in Category.objects.filter(parent__isnull=True)
    }
    created = []
    for type_name, category_name, subcategory_name in _read_rows(sheet_name):
        category_type = (
            Category.CategoryType.INCOME if type_name.lower() == "income" else Category.CategoryType.EXPENSE
        )
        parent = top_level.get(category_name.lower())
        if parent is None:
            parent = Category.objects.create(name=category_name, category_type=category_type)
            top_level[category_name.lower()] = parent
            created.append(category_name)
        if subcategory_name and not parent.subcategories.filter(name__iexact=subcategory_name).exists():
            Category.objects.create(name=subcategory_name, category_type=category_type, parent=parent)
            created.append(f"{category_name}: {subcategory_name}")
    return created
