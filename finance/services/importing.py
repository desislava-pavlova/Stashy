"""Importing a bank or finance-app file (CSV, XLSX, OFX or QFX) into one account.

`import_file()` is the one entry point. The whole file is read and checked before anything is
written, so a file Stashy can't read raises `ImportFileError` and leaves nothing behind.
"""

import csv
import hashlib
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.db import router, transaction

from audit.context import system_action

from .. import regional
from ..models import Account, Category, ImportBatch, Payee, Transaction, name_key
from ..money import parse_amount

SUPPORTED_EXTENSIONS = ("csv", "xlsx", "ofx", "qfx")
UNSUPPORTED_FILE_MESSAGE = "Unsupported file type. Please upload a .csv, .xlsx, .ofx, or .qfx file."


class ImportFileError(Exception):
    """The uploaded file couldn't be read at all - the message is shown to the user as is."""


@dataclass
class ImportResult:
    created: int = 0
    duplicates: int = 0
    unknown_account: int = 0
    unreadable: int = 0
    # The date format the rows were read with (CSV and XLSX only).
    date_format: str | None = None
    # Set when a CSV wasn't UTF-8 and was read as Windows-1251 instead.
    encoding_note: str = ""


def import_file(upload, account: Account, date_format: str = "auto") -> ImportResult:
    """Import every transaction in `upload` into `account`, as one ImportBatch.

    A row whose Account column names another account goes there instead. `date_format` is
    "auto" (read it off the file's own dates) or one of regional.DATE_FORMAT_CHOICES; OFX
    files carry real dates and ignore it.
    """
    extension = Path(upload.name).suffix.lower().lstrip(".")
    encoding_note = ""
    if extension == "csv":
        fieldnames, rows, encoding_note = _read_csv(upload)
    elif extension == "xlsx":
        fieldnames, rows = _read_xlsx(upload)
    elif extension in ("ofx", "qfx"):
        ofx_rows = _read_ofx(upload)
    else:
        raise ImportFileError(UNSUPPORTED_FILE_MESSAGE)

    with transaction.atomic(using=router.db_for_write(Transaction)), system_action(f"Imported from {upload.name}"):
        batch = ImportBatch.objects.create(account=account, source_filename=upload.name)
        if extension in ("csv", "xlsx"):
            result = _import_rows(fieldnames, rows, account, batch, date_format or "auto")
        else:
            result = _import_ofx_rows(ofx_rows, account, batch)
        batch.row_count = result.created
        batch.duplicate_count = result.duplicates
        batch.save(update_fields=["row_count", "duplicate_count"])
    result.encoding_note = encoding_note
    return result


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------


def _read_csv(upload):
    """(fieldnames, rows, encoding note) from an uploaded CSV.

    UTF-8 first (with or without Excel's BOM). Bulgarian bank exports are often Windows-1251,
    so that's the fallback (Claude's call, plan 02 item 1.10) - and the summary says so, since a
    Western file in Windows-1252 would come out with Cyrillic letters in place of its accents.
    """
    raw = upload.read()
    try:
        text, note = raw.decode("utf-8-sig"), ""
    except UnicodeDecodeError:
        try:
            text = raw.decode("cp1251")
        except UnicodeDecodeError:
            raise ImportFileError(
                "Stashy couldn't read this file's text. Save it as \"CSV UTF-8\" and try again."
            ) from None
        note = (
            "The file wasn't UTF-8, so it was read as Windows-1251 (Cyrillic) - if names look wrong, "
            'delete what was imported, save the file as "CSV UTF-8" and import it again.'
        )
    try:
        reader = csv.DictReader(io.StringIO(text, newline=""))
        rows = list(reader)
    except csv.Error as error:
        raise ImportFileError(f"This doesn't look like a CSV file Stashy can read ({error}).") from None
    return reader.fieldnames or [], rows, note


def _read_xlsx(upload):
    """(fieldnames, rows) from the first sheet of an uploaded Excel file."""
    from zipfile import BadZipFile

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(upload, data_only=True, read_only=True)
        all_rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError, TypeError):
        raise ImportFileError(
            "Stashy couldn't open this Excel file - it may be damaged or not really an .xlsx file."
        ) from None

    header = all_rows[0] if all_rows else []
    fieldnames = [str(cell).strip() if cell is not None else "" for cell in header]

    rows = []
    for values in all_rows[1:]:
        row = {}
        for name, value in zip(fieldnames, values, strict=False):
            if not name:
                continue
            if isinstance(value, datetime):
                value = value.date()
            if isinstance(value, date):
                row[name] = value.isoformat()
            elif value is None:
                row[name] = ""
            else:
                row[name] = str(value)
        rows.append(row)
    return fieldnames, rows


def _read_ofx(upload):
    """Every transaction in an OFX/QFX file as (date, raw amount, payee name, memo)."""
    from ofxparse import OfxParser

    try:
        ofx = OfxParser.parse(upload)
        return [
            (txn.date.date(), txn.amount, (txn.payee or "").strip(), getattr(txn, "memo", "") or "")
            for ofx_account in ofx.accounts
            for txn in ofx_account.statement.transactions
        ]
    except Exception:
        raise ImportFileError(
            "Stashy couldn't read this OFX/QFX file - it may be damaged or in a different format."
        ) from None


# ---------------------------------------------------------------------------
# Turning rows into transactions
# ---------------------------------------------------------------------------

_RECONCILED_FALSE_VALUES = {"", "n", "no", "false", "0", "unreconciled", "uncleared", "u"}


def _import_rows(fieldnames, rows, account, batch, date_format="auto"):
    """Shared row processing for CSV and XLSX imports. Recognizes Date, Num, Account, Payee,
    Category, Memo, Amount and Reconciled columns (the layout Quicken/Money "transaction
    report" exports use) - only Date needs to be present in a row; Num, Payee, Category, Memo
    and Reconciled can all be blank. Amount still has to parse, since a transaction can't be
    saved without one, and an Account column naming an account that doesn't exist is skipped
    rather than guessed at (D9 backlog "Accounts must already exist")."""
    # Normalize header names so common bank export variants all work.
    fieldmap = {name.strip().lower(): name for name in fieldnames if name}

    def get(row, *keys):
        for k in keys:
            if k in fieldmap:
                return (row.get(fieldmap[k]) or "").strip()
        return ""

    if date_format == "auto" or not date_format:
        detected = _detect_date_format(get(row, "date", "transaction date") for row in rows)
        date_format = detected or regional.current().date_format
    date_formats = regional.Preferences(date_format=date_format).date_input_formats

    accounts_by_name = {a.name.lower(): a for a in Account.objects.all()}

    result = ImportResult(date_format=date_format)
    for row in rows:
        if not any(isinstance(value, str) and value.strip() for value in row.values()):
            continue  # a completely empty line isn't a skipped transaction
        raw_date = get(row, "date", "transaction date")
        raw_amount = get(row, "amount")
        raw_num = get(row, "num", "number", "check number")
        payee_name = get(row, "payee", "description", "name")
        memo = get(row, "memo", "notes")
        raw_account = get(row, "account")
        raw_category = get(row, "category")
        raw_reconciled = get(row, "reconciled", "cleared", "status")

        txn_date = _parse_date(raw_date, date_formats) if raw_date else None
        amount = parse_amount(raw_amount) if raw_amount else None
        if txn_date is None or amount is None:
            result.unreadable += 1
            continue

        row_account = account
        if raw_account:
            row_account = accounts_by_name.get(raw_account.lower())
            if row_account is None:
                result.unknown_account += 1
                continue

        fp = _fingerprint(row_account.id, txn_date, amount, payee_name)
        if Transaction.objects.filter(import_fingerprint=fp).exists():
            result.duplicates += 1
            continue

        payee = Payee.objects.resolve(payee_name, reason=None) if payee_name else None
        category = _resolve_category(raw_category, amount) if raw_category else _guessed_category(payee)

        Transaction.objects.create(
            account=row_account,
            date=txn_date,
            transaction_number=_parse_transaction_number(raw_num),
            payee=payee,
            memo=memo,
            amount=amount,
            reconciled=raw_reconciled.strip().lower() not in _RECONCILED_FALSE_VALUES,
            import_batch=batch,
            import_fingerprint=fp,
            category=category,
        )
        result.created += 1

    return result


def _import_ofx_rows(ofx_rows, account, batch):
    result = ImportResult()
    for txn_date, raw_amount, payee_name, memo in ofx_rows:
        amount = parse_amount(raw_amount)
        if amount is None:
            result.unreadable += 1
            continue
        fp = _fingerprint(account.id, txn_date, amount, payee_name)
        if Transaction.objects.filter(import_fingerprint=fp).exists():
            result.duplicates += 1
            continue
        payee = Payee.objects.resolve(payee_name, reason=None) if payee_name else None
        Transaction.objects.create(
            account=account,
            date=txn_date,
            payee=payee,
            memo=memo,
            amount=amount,
            import_batch=batch,
            import_fingerprint=fp,
            category=_guessed_category(payee),
        )
        result.created += 1
    # OFX has no per-row account text field and no ambiguous textual dates to detect.
    return result


def _fingerprint(account_id, txn_date, amount, payee):
    """What makes two imported rows "the same transaction": importing a file twice skips
    every row whose fingerprint is already stored."""
    raw = f"{account_id}|{txn_date.isoformat()}|{amount:.2f}|{payee}".encode()
    return hashlib.sha256(raw).hexdigest()


def _guessed_category(payee):
    """The category to pre-fill for an imported row with no Category column value: this
    payee's most recently used category, if any."""
    if payee is None:
        return None
    history = payee.category_history()
    return history[0].category if history else None


_CATEGORY_SEPARATOR_RE = re.compile(r"\s*[:>]\s*")


def _resolve_category(raw, amount):
    """Resolve the category named in an import row, auto-creating it if Stashy doesn't already
    have one by that name (Phase 9) - the main real-world use of this importer is a large
    one-time historical migration, and requiring every category to exist beforehand would be
    exactly the kind of friction that makes migrations painful. Matching is case-insensitive,
    so "Groceries" from a file never creates a duplicate of an existing "groceries". Understands
    "Parent:Child" or "Parent > Child" (Quicken's and this app's own separators) for one level
    of subcategory - a third level, if a file somehow has one, folds into the child's own name
    rather than erroring, since this app doesn't support deeper nesting anyway.
    """
    raw = (raw or "").strip()
    if not raw:
        return None
    parts = _CATEGORY_SEPARATOR_RE.split(raw, maxsplit=1)
    parent_name = parts[0].strip()
    child_name = parts[1].strip() if len(parts) > 1 else ""
    if not parent_name:
        return None

    category_type = Category.CategoryType.INCOME if amount > 0 else Category.CategoryType.EXPENSE
    parent = Category.objects.filter(name_key=name_key(parent_name), parent__isnull=True, account__isnull=True).first()
    if parent is None:
        parent = Category.objects.create(name=parent_name, category_type=category_type, is_auto_generated=True)

    if not child_name:
        return parent

    child = Category.objects.filter(name_key=name_key(child_name), parent=parent).first()
    if child is None:
        # A new subcategory always takes its parent's type - a mismatched pair (an "income"
        # category with an "expense" child) would only confuse reports for no real benefit.
        child = Category.objects.create(
            name=child_name, category_type=parent.category_type, parent=parent, is_auto_generated=True
        )
    return child


# A day value over 12 in either position settles day-first vs month-first; a 4-digit first
# component is ISO. Ambiguous rows (both parts <= 12) are skipped - the caller falls back to
# the profile's own format if nothing in the whole file was ever unambiguous.
_DATE_SNIFF_RE = re.compile(r"^(\d{1,4})[./-](\d{1,2})[./-](\d{1,4})$")


def _detect_date_format(raw_dates):
    for raw in raw_dates:
        match = _DATE_SNIFF_RE.match((raw or "").strip())
        if not match:
            continue
        a, b, _c = match.groups()
        if len(a) == 4:
            return "ymd"
        if int(a) > 12:
            return "dmy"
        if int(b) > 12:
            return "mdy"
    return None


def _parse_date(raw, formats):
    for fmt in formats:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


# The largest value Django guarantees a PositiveIntegerField can hold on every database.
_MAX_TRANSACTION_NUMBER = 2_147_483_647


def _parse_transaction_number(raw):
    """A Num cell as a transaction number, or None ("no number") when it isn't a whole number
    Stashy can store - blank, negative, huge, "inf", "12.5" or plain text."""
    try:
        number = Decimal(raw.strip())
    except (InvalidOperation, AttributeError):
        return None
    if not number.is_finite() or number != number.to_integral_value():
        return None
    number = int(number)
    return number if 0 <= number <= _MAX_TRANSACTION_NUMBER else None
